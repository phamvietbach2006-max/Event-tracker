import json
import re
import sys
import unicodedata
import openpyxl

EXCEL_PATH = r"C:\Users\phamv\Downloads\[ĐH-SEEE] PHÂN KÍP ĐẠI HỘI CHI ĐOÀN - CHI HỘI NĂM HỌC 2026 - 2027.xlsx"
SHEET_NAME = "Ngày 1810 (Off)"

# K71: Đại hội thành lập (4 mục hồ sơ)
TASKS_K71 = [
    ["1. Phương hướng", "todo"],
    ["2. Nghị quyết Đại hội", "todo"],
    ["3. Biên bản Đại hội", "todo"],
    ["4. Danh sách Đoàn viên và Cán bộ lớp", "todo"],
]

# K70: Đại hội nhiệm kỳ (8 mục hồ sơ)
TASKS_K70 = [
    ["1. Biên bản ĐHCĐ", "todo"],
    ["2. Biên bản kiểm phiếu bầu BCH CĐ", "todo"],
    ["3. Biên bản họp phân công BCH CĐ", "todo"],
    ["4. Bản kiểm điểm BCH", "todo"],
    ["5. Nghị quyết Đại hội", "todo"],
    ["6. Báo cáo tổng kết", "todo"],
    ["7. Phương hướng nhiệm kỳ mới", "todo"],
    ["8. Danh sách Đoàn viên và Cán bộ lớp", "todo"],
]


def slug_email(name: str) -> str:
    s = unicodedata.normalize("NFD", name)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.replace("đ", "d").replace("Đ", "d").lower()
    s = re.sub(r"[^a-z0-9]+", ".", s).strip(".")
    return f"{s}@seee.hust.edu.vn"


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb[SHEET_NAME]

    users_map = {
        "admin@example.com": {
            "email": "admin@example.com",
            "name": "Quản trị viên (Ban TC-KT)",
            "isAdmin": True,
        },
        "nguyen.long.vu@seee.hust.edu.vn": {
            "email": "nguyen.long.vu@seee.hust.edu.vn",
            "name": "Nguyễn Long Vũ",
            "isAdmin": True,
        },
        "truong.tuan.kiet@seee.hust.edu.vn": {
            "email": "truong.tuan.kiet@seee.hust.edu.vn",
            "name": "Trương Tuấn Kiệt",
            "isAdmin": True,
        },
    }

    unit_name = "Trường Điện-Điện tử"
    floors = {
        "1": {"name": f"[{unit_name}] Tầng 1 tòa D9", "rooms": []},
        "2": {"name": f"[{unit_name}] Tầng 2 tòa D9", "rooms": []},
        "3": {"name": f"[{unit_name}] Tầng 3 tòa D9", "rooms": []},
        "4": {"name": f"[{unit_name}] Tầng 4 tòa D9", "rooms": []},
        "5": {"name": f"[{unit_name}] Tầng 5 tòa D9", "rooms": []},
    }

    current_kip = ""
    for r in range(7, ws.max_row + 1):
        kip_val = ws.cell(r, 3).value
        if kip_val and "Kíp" in str(kip_val):
            lines = [x.strip() for x in str(kip_val).splitlines() if x.strip()]
            current_kip = f"{lines[0]} ({lines[1]})" if len(lines) > 1 else lines[0]

        stt = ws.cell(r, 1).value
        room_code = ws.cell(r, 4).value
        cls_name = ws.cell(r, 5).value

        if isinstance(stt, (int, float)) and room_code and cls_name:
            room_code = str(room_code).strip()
            cls_name = str(cls_name).strip()
            floor_num = room_code.split("-")[1][0]
            is_k71 = "K71" in cls_name

            owners = []
            for c in range(9, 21, 2):
                person = ws.cell(r, c).value
                if person and str(person).strip():
                    p_name = str(person).strip()
                    em = slug_email(p_name)
                    if em not in users_map:
                        users_map[em] = {"email": em, "name": p_name, "isAdmin": True}
                    if em not in owners:
                        owners.append(em)

            room_title = f"[{current_kip} · {room_code}] {cls_name}"
            floors[floor_num]["rooms"].append(
                {
                    "name": room_title,
                    "status": "red",
                    "note": "",
                    "owners": owners,
                    "tasks": [list(t) for t in (TASKS_K71 if is_k71 else TASKS_K70)],
                }
            )

    seed = {
        "users": list(users_map.values()),
        "groups": [floors[k] for k in ["1", "2", "3", "4", "5"]],
    }

    with open("seed.json", "w", encoding="utf-8") as f:
        json.dump(seed, f, ensure_ascii=False, indent=2)

    print("Đã xuất seed.json thành công:")
    total_rooms = 0
    for g in seed["groups"]:
        c = len(g["rooms"])
        total_rooms += c
        print(f"  - {g['name']}: {c} lớp")
    print(f"Tổng cộng: {total_rooms} lớp, {len(seed['users'])} tài khoản.")


if __name__ == "__main__":
    main()
